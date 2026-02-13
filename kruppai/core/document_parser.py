"""Universal document parser for construction documents.

Extracts text and structured data from PDF, DOCX, XLSX, images, and plain text.
Returns a uniform ParsedDocument dataclass regardless of source format.
"""

import base64
import mimetypes
from dataclasses import dataclass, field
from pathlib import Path

SUPPORTED_EXTENSIONS: dict[str, str] = {
    ".pdf": "pdf",
    ".docx": "docx",
    ".xlsx": "xlsx",
    ".xls": "xlsx",
    ".png": "image",
    ".jpg": "image",
    ".jpeg": "image",
    ".gif": "image",
    ".bmp": "image",
    ".tiff": "image",
    ".tif": "image",
    ".webp": "image",
    ".txt": "text",
    ".md": "text",
    ".csv": "text",
}


@dataclass
class ParsedDocument:
    """Uniform parsed document representation."""

    source_path: Path
    file_type: str
    text: str
    tables: list[list[list[str]]] = field(default_factory=list)
    images: list[dict] = field(default_factory=list)
    metadata: dict = field(default_factory=dict)
    page_count: int = 0
    char_count: int = 0
    error: str | None = None


def parse_file(file_path: Path, max_chars: int = 100_000) -> ParsedDocument:
    """Parse any supported file into a ParsedDocument.

    Args:
        file_path: Path to the file to parse.
        max_chars: Maximum characters to extract (prevents API cost blowup).

    Returns:
        ParsedDocument with extracted content or error message.
    """
    if not file_path.exists():
        return ParsedDocument(
            source_path=file_path,
            file_type="unknown",
            text="",
            error=f"File not found: {file_path}",
        )

    ext = file_path.suffix.lower()
    file_type = SUPPORTED_EXTENSIONS.get(ext)

    if file_type is None:
        return ParsedDocument(
            source_path=file_path,
            file_type="unsupported",
            text="",
            error=f"Unsupported file format: {ext}",
        )

    parsers = {
        "pdf": parse_pdf,
        "docx": parse_docx,
        "xlsx": parse_xlsx,
        "image": parse_image,
        "text": parse_text,
    }

    return parsers[file_type](file_path, max_chars)


def parse_pdf(file_path: Path, max_chars: int = 100_000) -> ParsedDocument:
    """Extract text and tables from a PDF using PyMuPDF."""
    try:
        import fitz  # PyMuPDF
    except ImportError:
        return ParsedDocument(
            source_path=file_path,
            file_type="pdf",
            text="",
            error="PyMuPDF not installed. Install with: pip install PyMuPDF",
        )

    try:
        doc = fitz.open(str(file_path))
    except Exception as e:
        return ParsedDocument(
            source_path=file_path,
            file_type="pdf",
            text="",
            error=f"Failed to open PDF: {e}",
        )

    text_parts: list[str] = []
    tables: list[list[list[str]]] = []
    total_chars = 0

    try:
        for page_num, page in enumerate(doc):
            page_text = page.get_text()
            if total_chars + len(page_text) > max_chars:
                remaining = max_chars - total_chars
                if remaining > 0:
                    text_parts.append(page_text[:remaining])
                text_parts.append("\n[TRUNCATED — document exceeds character limit]")
                break
            text_parts.append(page_text)
            total_chars += len(page_text)

            # Extract tables if available
            try:
                page_tables = page.find_tables()
                for table in page_tables:
                    extracted = table.extract()
                    if extracted:
                        tables.append(
                            [[str(cell) if cell else "" for cell in row] for row in extracted]
                        )
            except Exception:
                pass  # Table extraction is best-effort

        full_text = "\n".join(text_parts)
        metadata = {
            "pages": doc.page_count,
            "author": doc.metadata.get("author", ""),
            "title": doc.metadata.get("title", ""),
            "created": doc.metadata.get("creationDate", ""),
        }

        return ParsedDocument(
            source_path=file_path,
            file_type="pdf",
            text=full_text,
            tables=tables,
            metadata=metadata,
            page_count=doc.page_count,
            char_count=len(full_text),
        )
    finally:
        doc.close()


def parse_docx(file_path: Path, max_chars: int = 100_000) -> ParsedDocument:
    """Extract text and tables from a DOCX file."""
    try:
        from docx import Document
    except ImportError:
        return ParsedDocument(
            source_path=file_path,
            file_type="docx",
            text="",
            error="python-docx not installed. Install with: pip install python-docx",
        )

    try:
        doc = Document(str(file_path))
    except Exception as e:
        return ParsedDocument(
            source_path=file_path,
            file_type="docx",
            text="",
            error=f"Failed to open DOCX: {e}",
        )

    text_parts: list[str] = []
    total_chars = 0

    for para in doc.paragraphs:
        if total_chars + len(para.text) > max_chars:
            text_parts.append("\n[TRUNCATED — document exceeds character limit]")
            break
        text_parts.append(para.text)
        total_chars += len(para.text)

    tables: list[list[list[str]]] = []
    for table in doc.tables:
        rows = []
        for row in table.rows:
            rows.append([cell.text.strip() for cell in row.cells])
        tables.append(rows)

    full_text = "\n".join(text_parts)

    metadata: dict = {}
    if doc.core_properties:
        metadata["author"] = doc.core_properties.author or ""
        metadata["title"] = doc.core_properties.title or ""

    return ParsedDocument(
        source_path=file_path,
        file_type="docx",
        text=full_text,
        tables=tables,
        metadata=metadata,
        page_count=0,  # DOCX doesn't expose page count without rendering
        char_count=len(full_text),
    )


def parse_xlsx(file_path: Path, max_chars: int = 100_000) -> ParsedDocument:
    """Extract cell values from all sheets of an XLSX file."""
    try:
        from openpyxl import load_workbook
    except ImportError:
        return ParsedDocument(
            source_path=file_path,
            file_type="xlsx",
            text="",
            error="openpyxl not installed. Install with: pip install openpyxl",
        )

    try:
        wb = load_workbook(str(file_path), read_only=True, data_only=True)
    except Exception as e:
        return ParsedDocument(
            source_path=file_path,
            file_type="xlsx",
            text="",
            error=f"Failed to open XLSX: {e}",
        )

    text_parts: list[str] = []
    tables: list[list[list[str]]] = []
    total_chars = 0

    try:
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            text_parts.append(f"\n--- Sheet: {sheet_name} ---\n")
            sheet_rows: list[list[str]] = []

            for row in ws.iter_rows():
                cell_values = [
                    str(cell.value) if cell.value is not None else ""
                    for cell in row
                ]
                row_text = " | ".join(cell_values)

                if total_chars + len(row_text) > max_chars:
                    text_parts.append(
                        "\n[TRUNCATED — document exceeds character limit]"
                    )
                    break
                text_parts.append(row_text)
                total_chars += len(row_text)
                sheet_rows.append(cell_values)

            if sheet_rows:
                tables.append(sheet_rows)
    finally:
        wb.close()

    full_text = "\n".join(text_parts)
    metadata = {"sheets": wb.sheetnames}

    return ParsedDocument(
        source_path=file_path,
        file_type="xlsx",
        text=full_text,
        tables=tables,
        metadata=metadata,
        page_count=len(wb.sheetnames),
        char_count=len(full_text),
    )


def parse_image(file_path: Path, max_chars: int = 100_000) -> ParsedDocument:
    """Encode an image as base64 for Claude vision API."""
    try:
        data = file_path.read_bytes()
    except Exception as e:
        return ParsedDocument(
            source_path=file_path,
            file_type="image",
            text="",
            error=f"Failed to read image: {e}",
        )

    encoded = base64.b64encode(data).decode("utf-8")
    mime_type = mimetypes.guess_type(str(file_path))[0] or "image/png"

    return ParsedDocument(
        source_path=file_path,
        file_type="image",
        text=f"[Image: {file_path.name}]",
        images=[
            {
                "base64": encoded,
                "media_type": mime_type,
                "description": file_path.name,
            }
        ],
        metadata={"file_size_bytes": len(data)},
        page_count=1,
        char_count=0,
    )


def parse_text(file_path: Path, max_chars: int = 100_000) -> ParsedDocument:
    """Read a plain text or markdown file."""
    try:
        text = file_path.read_text(encoding="utf-8")
    except Exception as e:
        return ParsedDocument(
            source_path=file_path,
            file_type="text",
            text="",
            error=f"Failed to read text file: {e}",
        )

    if len(text) > max_chars:
        text = text[:max_chars] + "\n[TRUNCATED — document exceeds character limit]"

    return ParsedDocument(
        source_path=file_path,
        file_type="text",
        text=text,
        metadata={},
        page_count=1,
        char_count=len(text),
    )
