"""Shared test fixtures for KruppAI test suite.

Provides:
- tmp_settings: Settings pointing to temp directories
- test_db: Initialized SQLite database in temp directory
- seeded_db: Database with realistic construction project data
- mock_api_client: API client returning canned responses
- sample_project: Realistic construction project dict
"""

import sqlite3
from pathlib import Path
from unittest.mock import MagicMock

import pytest

from kruppai.core.api_client import AnthropicClient, ApiResponse
from kruppai.core.config import Settings
from kruppai.core.context_manager import ContextManager
from kruppai.core.database import get_db, init_db
from kruppai.core.knowledge_base import KnowledgeBase
from kruppai.core.output_formatter import OutputFormatter


@pytest.fixture
def tmp_settings(tmp_path: Path) -> Settings:
    """Settings with all paths in temp directory."""
    return Settings(
        anthropic_api_key="sk-ant-test-key-not-real",
        db_path=tmp_path / "test.db",
        output_dir=tmp_path / "output",
        knowledge_dir=tmp_path / "knowledge",
        daily_cost_limit_cents=10000,
        monthly_cost_limit_cents=50000,
        _env_file=None,
    )


@pytest.fixture
def test_db(tmp_settings: Settings) -> Settings:
    """Empty initialized database. Returns the settings object."""
    init_db(tmp_settings)
    return tmp_settings


@pytest.fixture
def seeded_db(test_db: Settings) -> Settings:
    """Database pre-loaded with realistic construction project data."""
    settings = test_db
    with get_db(settings) as conn:
        # Company
        conn.execute(
            "INSERT INTO company "
            "(name, legal_name, address_line1, city, state, zip, phone, email, "
            "founded_year, description, specialties, service_area) "
            "VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)",
            (
                "Krupp General Contractors",
                "Krupp General Contractors, LLC",
                "1200 Construction Way",
                "Denver",
                "CO",
                "80202",
                "(303) 555-0100",
                "info@kruppgc.com",
                2005,
                "Full-service general contractor specializing in commercial and healthcare construction.",
                '["commercial", "healthcare", "education"]',
                "Front Range Colorado",
            ),
        )

        # Team members
        team_data = [
            ("Sarah", "Chen", "schen@kruppgc.com", "project_manager",
             "Senior Project Manager", '["PMP", "LEED AP"]', 15),
            ("Mike", "Rodriguez", "mrodriguez@kruppgc.com", "superintendent",
             "Senior Superintendent", '["OSHA 30", "First Aid/CPR"]', 22),
            ("James", "Park", "jpark@kruppgc.com", "project_engineer",
             "Project Engineer", '["EIT"]', 4),
        ]
        for first, last, email, role, title, certs, yrs in team_data:
            conn.execute(
                "INSERT INTO team_members "
                "(first_name, last_name, email, role, title, certifications, "
                "years_experience) "
                "VALUES (?, ?, ?, ?, ?, ?, ?)",
                (first, last, email, role, title, certs, yrs),
            )

        # Project
        conn.execute(
            "INSERT INTO projects "
            "(project_code, name, client_name, client_contact_name, "
            "address, city, state, zip, latitude, longitude, "
            "original_contract_cents, current_contract_cents, "
            "notice_to_proceed_date, substantial_completion_date, "
            "current_percent_complete, project_type, delivery_method, "
            "contract_type, square_footage, number_of_floors, status) "
            "VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)",
            (
                "KRUPP-2026-TEST",
                "City Center Medical Office Building",
                "Regional Health Partners",
                "Dr. Amanda Foster",
                "500 Main Street",
                "Denver",
                "CO",
                "80202",
                39.7392,
                -104.9903,
                1_250_000_000,  # $12.5M
                1_285_000_000,  # $12.85M (with approved COs)
                "2025-09-01",
                "2026-09-15",
                35.0,
                "healthcare",
                "negotiated",
                "gmp",
                45000,
                3,
                "active",
            ),
        )

        # Assign team to project
        conn.execute(
            "INSERT INTO project_team (project_id, team_member_id, project_role, is_primary) "
            "VALUES (1, 1, 'project_manager', 1)"
        )
        conn.execute(
            "INSERT INTO project_team (project_id, team_member_id, project_role, is_primary) "
            "VALUES (1, 2, 'superintendent', 1)"
        )
        conn.execute(
            "INSERT INTO project_team (project_id, team_member_id, project_role, is_primary) "
            "VALUES (1, 3, 'project_engineer', 1)"
        )

        # Subcontractors
        conn.execute(
            "INSERT INTO subcontractors "
            "(company_name, contact_name, email, phone, trade, csi_division, rating) "
            "VALUES (?, ?, ?, ?, ?, ?, ?)",
            (
                "ABC Electrical Services",
                "Tom Baker",
                "tbaker@abcelectric.com",
                "(303) 555-0200",
                "electrical",
                "26",
                4.2,
            ),
        )
        conn.execute(
            "INSERT INTO subcontractors "
            "(company_name, contact_name, email, phone, trade, csi_division, rating) "
            "VALUES (?, ?, ?, ?, ?, ?, ?)",
            (
                "Pacific Mechanical",
                "Lisa Nguyen",
                "lnguyen@pacmech.com",
                "(303) 555-0300",
                "mechanical",
                "23",
                4.5,
            ),
        )

        # Link subs to project
        conn.execute(
            "INSERT INTO project_subcontractors "
            "(project_id, subcontractor_id, contract_value_cents, "
            "scope_description, contract_status) "
            "VALUES (1, 1, ?, ?, ?)",
            (120_000_000, "Complete electrical systems including power, lighting, fire alarm", "executed"),
        )
        conn.execute(
            "INSERT INTO project_subcontractors "
            "(project_id, subcontractor_id, contract_value_cents, "
            "scope_description, contract_status) "
            "VALUES (1, 2, ?, ?, ?)",
            (180_000_000, "HVAC, plumbing, medical gas systems", "executed"),
        )

        # Action items
        action_items = [
            ("Follow up on RFI #12 response from architect", "Sarah Chen", "2026-02-15", "high"),
            ("Submit electrical panel shop drawings", "ABC Electrical Services", "2026-02-20", "normal"),
            ("Review updated project schedule", "Mike Rodriguez", "2026-02-18", "normal"),
            ("Provide revised HVAC ductwork layout", "Pacific Mechanical", "2026-02-22", "urgent"),
            ("Update owner on steel delivery delay", "Sarah Chen", "2026-02-14", "high"),
        ]
        for desc, assigned, due, priority in action_items:
            conn.execute(
                "INSERT INTO action_items "
                "(project_id, description, assigned_to, due_date, priority, status, source_skill) "
                "VALUES (1, ?, ?, ?, ?, 'open', 'meeting_minutes')",
                (desc, assigned, due, priority),
            )

    return settings


@pytest.fixture
def mock_api_response() -> ApiResponse:
    """A realistic mock API response."""
    return ApiResponse(
        content='{"report_date": "2026-02-13", "work_performed": "Concrete pour completed on 3rd floor."}',
        model="claude-sonnet-4-5-20250929",
        input_tokens=1500,
        output_tokens=800,
        cost_cents=2,
        duration_ms=3500,
    )


@pytest.fixture
def mock_api_client(mock_api_response: ApiResponse) -> AnthropicClient:
    """API client that returns predictable responses without network calls."""
    client = MagicMock(spec=AnthropicClient)
    client.call.return_value = mock_api_response
    client.settings = Settings(
        anthropic_api_key="sk-ant-test-key",
        _env_file=None,
    )
    return client


@pytest.fixture
def sample_project() -> dict:
    """Realistic construction project dict."""
    return {
        "id": 1,
        "project_code": "KRUPP-2026-TEST",
        "name": "City Center Medical Office Building",
        "client_name": "Regional Health Partners",
        "client_contact_name": "Dr. Amanda Foster",
        "address": "500 Main Street",
        "city": "Denver",
        "state": "CO",
        "zip": "80202",
        "latitude": 39.7392,
        "longitude": -104.9903,
        "original_contract_cents": 1_250_000_000,
        "current_contract_cents": 1_285_000_000,
        "current_percent_complete": 35.0,
        "project_type": "healthcare",
        "delivery_method": "negotiated",
        "contract_type": "gmp",
        "square_footage": 45000,
        "number_of_floors": 3,
        "status": "active",
    }


@pytest.fixture
def knowledge_dir(tmp_path: Path) -> Path:
    """Create a temp knowledge directory with sample files."""
    kd = tmp_path / "knowledge"
    kd.mkdir()
    (kd / "company_profile.md").write_text("# Krupp GC\nTest company profile.")
    (kd / "writing_standards.md").write_text("# Writing Standards\nBe professional.")
    (kd / "safety_standards.md").write_text("# Safety\nSafety first.")
    return kd


@pytest.fixture
def sample_docx(tmp_path: Path) -> Path:
    """Create a minimal DOCX test file."""
    from docx import Document

    doc = Document()
    doc.add_paragraph("Test document content for KruppAI testing.")
    table = doc.add_table(rows=2, cols=3)
    table.rows[0].cells[0].text = "Header A"
    table.rows[0].cells[1].text = "Header B"
    table.rows[0].cells[2].text = "Header C"
    table.rows[1].cells[0].text = "Value 1"
    table.rows[1].cells[1].text = "Value 2"
    table.rows[1].cells[2].text = "Value 3"
    path = tmp_path / "test_doc.docx"
    doc.save(str(path))
    return path


@pytest.fixture
def sample_xlsx(tmp_path: Path) -> Path:
    """Create a minimal XLSX test file with 2 sheets."""
    from openpyxl import Workbook

    wb = Workbook()
    ws1 = wb.active
    ws1.title = "Estimate"
    ws1.append(["CSI Code", "Description", "Amount"])
    ws1.append(["03", "Concrete", "450000"])
    ws1.append(["26", "Electrical", "1200000"])

    ws2 = wb.create_sheet("Summary")
    ws2.append(["Category", "Total"])
    ws2.append(["Total", "1650000"])

    path = tmp_path / "test_spreadsheet.xlsx"
    wb.save(str(path))
    return path


@pytest.fixture
def sample_image(tmp_path: Path) -> Path:
    """Create a minimal PNG test image."""
    from PIL import Image

    img = Image.new("RGB", (100, 100), color="red")
    path = tmp_path / "test_photo.png"
    img.save(str(path))
    return path
