"""Tests for the Toolbox Safety Talk Generator skill."""

import json
from datetime import date
from unittest.mock import MagicMock

import pytest

from kruppai.core.api_client import ApiResponse
from kruppai.core.config import Settings
from kruppai.core.context_manager import ContextManager
from kruppai.core.database import get_db
from kruppai.core.knowledge_base import KnowledgeBase
from kruppai.core.output_formatter import OutputFormatter
from kruppai.skills.safety_talk import SafetyTalkSkill


MOCK_SAFETY_TALK_TEXT = """## Fall Protection — Steel Erection Safety

## Introduction

Good morning, crew. Next week we begin steel erection on the penthouse structure, and that means we're going to be working at heights consistently for the next several weeks. Falls from height remain the number one cause of death in construction, so let's take a few minutes to review our fall protection procedures.

## Key Hazards

Working at heights during steel erection presents several specific hazards:
- Unprotected leading edges as new steel is placed
- Openings in floor and roof decking
- Unstable walking surfaces on open steel beams
- Falling objects from overhead work
- Weather conditions (wind, rain, ice) making surfaces slippery

## Protective Measures

1. **100% tie-off** is required when working above 6 feet — no exceptions.
2. Inspect your harness before each use: check webbing for cuts, frayed edges, or burns. Check D-rings and buckles for deformation.
3. Use retractable self-retracting lifelines (SRLs) whenever possible — they limit free fall distance.
4. Never use a body belt for fall arrest — full body harness only.
5. Secure all tools and materials to prevent dropped objects.
6. Report any damaged fall protection equipment immediately.

## Real-World Example

Last year on a similar project, an ironworker stepped backward off an unprotected edge at 20 feet. His harness and SRL arrested the fall within 2 feet. Without that equipment, it would have been a 20-foot fall onto concrete. That harness saved his life and he was back to work the next day. The equipment works — but only if you wear it.

## Discussion Questions

1. What's the most challenging part of maintaining 100% tie-off during steel erection? How can we plan our anchor points better?
2. When was the last time you inspected your harness? What did you look for?
3. Have you ever seen a near-miss involving a fall? What happened and what could have been done differently?

## Summary

- Falls are the #1 killer in construction — take fall protection seriously
- 100% tie-off above 6 feet, no exceptions
- Inspect your harness every single day before use
- Use SRLs when available for the shortest possible fall distance
- Report damaged equipment — never use compromised gear

## OSHA References

- 29 CFR 1926.501 — Fall Protection Scope/Requirements
- 29 CFR 1926.502 — Fall Protection Systems Criteria
- 29 CFR 1926.760 — Fall Protection for Steel Erection
- 29 CFR 1926.1053 — Ladders
"""


def _make_skill(settings: Settings) -> SafetyTalkSkill:
    """Build a SafetyTalkSkill with mocked API client."""
    mock_api = MagicMock()
    mock_api.call.return_value = ApiResponse(
        content=MOCK_SAFETY_TALK_TEXT,
        model="claude-sonnet-4-5-20250929",
        input_tokens=800,
        output_tokens=1100,
        cost_cents=2,
        duration_ms=3500,
    )
    mock_api.settings = settings
    return SafetyTalkSkill(
        settings=settings,
        api_client=mock_api,
        formatter=OutputFormatter(settings),
        context_manager=ContextManager(settings),
        knowledge_base=KnowledgeBase(settings),
    )


class TestSafetyTalkValidation:
    """Test input validation for safety talks."""

    def test_validate_input_valid(self, seeded_db: Settings) -> None:
        skill = _make_skill(seeded_db)
        result = skill.validate_input(
            topic="Working at heights - fall protection review",
        )
        assert result["topic"] == "Working at heights - fall protection review"
        assert result["project_id"] is None

    def test_validate_input_with_project(self, seeded_db: Settings) -> None:
        skill = _make_skill(seeded_db)
        result = skill.validate_input(
            topic="Working at heights - fall protection",
            project="KRUPP-2026-TEST",
        )
        assert result["project_id"] == 1
        assert result["project_code"] == "KRUPP-2026-TEST"

    def test_validate_input_missing_topic(self, seeded_db: Settings) -> None:
        skill = _make_skill(seeded_db)
        with pytest.raises(ValueError, match="Safety topic is required"):
            skill.validate_input()

    def test_validate_input_topic_too_short(self, seeded_db: Settings) -> None:
        skill = _make_skill(seeded_db)
        with pytest.raises(ValueError, match="at least 5 characters"):
            skill.validate_input(topic="Hi")

    def test_validate_input_invalid_season(self, seeded_db: Settings) -> None:
        skill = _make_skill(seeded_db)
        with pytest.raises(ValueError, match="Invalid season"):
            skill.validate_input(
                topic="Heat stress awareness for the crew",
                season="monsoon",
            )

    def test_validate_input_with_season_and_trades(
        self, seeded_db: Settings
    ) -> None:
        skill = _make_skill(seeded_db)
        result = skill.validate_input(
            topic="Heat stress awareness for the crew",
            season="summer",
            trades="ironworkers, concrete finishers",
        )
        assert result["season"] == "summer"
        assert len(result["trades"]) == 2
        assert "ironworkers" in result["trades"]


class TestSafetyTalkPrompt:
    """Test prompt building for safety talks."""

    def test_build_prompt_structure(self, seeded_db: Settings) -> None:
        skill = _make_skill(seeded_db)
        validated = {
            "project_id": None,
            "project_code": None,
            "topic": "Fall protection review",
            "season": None,
            "trades": [],
        }
        messages = skill.build_prompt(validated, None)
        assert isinstance(messages, list)
        assert len(messages) >= 1
        assert messages[0]["role"] == "user"
        assert "SAFETY TOPIC" in messages[0]["content"]

    def test_build_prompt_includes_context(self, seeded_db: Settings) -> None:
        skill = _make_skill(seeded_db)
        context = skill.context_manager.load_by_code("KRUPP-2026-TEST")
        validated = {
            "project_id": 1,
            "project_code": "KRUPP-2026-TEST",
            "topic": "Fall protection review",
            "season": "winter",
            "trades": ["ironworkers"],
        }
        messages = skill.build_prompt(validated, context)
        content = messages[0]["content"]
        assert "Fall protection" in content


class TestSafetyTalkOutput:
    """Test output formatting for safety talks."""

    def test_format_output_creates_file(self, seeded_db: Settings) -> None:
        skill = _make_skill(seeded_db)
        validated = {
            "project_id": None,
            "project_code": None,
            "topic": "Fall protection review",
            "season": None,
            "trades": [],
        }
        output_path = skill.format_output(
            MOCK_SAFETY_TALK_TEXT, None, validated
        )
        assert output_path.exists()
        assert output_path.suffix == ".docx"

    def test_format_output_content(self, seeded_db: Settings) -> None:
        skill = _make_skill(seeded_db)
        validated = {
            "project_id": None,
            "project_code": None,
            "topic": "Fall protection review",
            "season": None,
            "trades": [],
        }
        output_path = skill.format_output(
            MOCK_SAFETY_TALK_TEXT, None, validated
        )

        from docx import Document

        doc = Document(str(output_path))
        full_text = "\n".join(p.text for p in doc.paragraphs)
        assert "Toolbox Safety Talk" in full_text

    def test_format_output_has_sign_in_sheet(self, seeded_db: Settings) -> None:
        skill = _make_skill(seeded_db)
        validated = {
            "project_id": None,
            "project_code": None,
            "topic": "Fall protection review",
            "season": None,
            "trades": [],
        }
        output_path = skill.format_output(
            MOCK_SAFETY_TALK_TEXT, None, validated
        )

        from docx import Document

        doc = Document(str(output_path))
        full_text = "\n".join(p.text for p in doc.paragraphs)
        assert "Attendee Sign-In" in full_text

    def test_format_output_persists_to_db(self, seeded_db: Settings) -> None:
        skill = _make_skill(seeded_db)
        validated = {
            "project_id": None,
            "project_code": None,
            "topic": "Fall protection review",
            "season": "winter",
            "trades": ["ironworkers"],
        }
        skill.format_output(MOCK_SAFETY_TALK_TEXT, None, validated)

        with get_db(seeded_db) as conn:
            cursor = conn.execute("SELECT * FROM safety_talks")
            row = cursor.fetchone()
            assert row is not None
            assert row["topic"] == "Fall protection review"
            assert row["season"] == "winter"


class TestSafetyTalkEndToEnd:
    """Test full safety talk lifecycle."""

    def test_execute_end_to_end_without_project(
        self, seeded_db: Settings
    ) -> None:
        skill = _make_skill(seeded_db)
        result = skill.execute(
            topic="Working at heights - fall protection for steel erection",
        )
        assert result.output_path.exists()
        assert result.skill_name == "safety_talk"
        assert result.project_code is None

    def test_execute_with_project(self, seeded_db: Settings) -> None:
        skill = _make_skill(seeded_db)
        result = skill.execute(
            topic="Working at heights - fall protection for steel erection",
            project="KRUPP-2026-TEST",
        )
        assert result.output_path.exists()
        assert result.project_code == "KRUPP-2026-TEST"
