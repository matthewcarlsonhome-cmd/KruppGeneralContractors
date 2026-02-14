"""Tests for the RFI Generator skill."""

import json
from datetime import date
from unittest.mock import MagicMock

import pytest

from kruppai.core.api_client import ApiResponse
from kruppai.core.config import Settings
from kruppai.core.context_manager import ContextManager
from kruppai.core.database import get_db
from kruppai.core.knowledge_base import KnowledgeBase
from kruppai.core.output_formatter import OutputFormatter
from kruppai.skills.rfi_generator import RfiGeneratorSkill


MOCK_RFI_JSON = json.dumps({
    "subject": "Structural Steel Beam Size Discrepancy at Grid Line C",
    "question": "The structural drawings (Sheet S-301) indicate W12x26 beams "
    "at grid line C, while Specification Section 05 12 00 calls for W14x30 "
    "at this location. Please clarify which beam size is correct for "
    "grid line C, as steel fabrication is scheduled to begin next week.",
    "spec_reference": "Section 05 12 00 - Structural Steel",
    "drawing_reference": "Sheet S-301",
    "cost_impact": "potential",
    "cost_impact_notes": "W14x30 beams are heavier and more expensive than W12x26. "
    "If the larger size is correct, there will be a cost increase for the steel package.",
    "schedule_impact": "potential",
    "schedule_impact_notes": "Steel fabrication begins next week. A delayed response "
    "could impact the fabrication schedule and subsequent erection dates.",
    "suggested_solution": "Based on the structural loading requirements and the "
    "specification language, the W14x30 appears to be the intended design. "
    "Recommend confirming via revised structural drawings.",
    "priority": "high",
})


def _make_skill(settings: Settings) -> RfiGeneratorSkill:
    """Build an RfiGeneratorSkill with mocked API client."""
    mock_api = MagicMock()
    mock_api.call.return_value = ApiResponse(
        content=MOCK_RFI_JSON,
        model="claude-sonnet-4-5-20250929",
        input_tokens=1200,
        output_tokens=700,
        cost_cents=1,
        duration_ms=2800,
    )
    mock_api.settings = settings
    return RfiGeneratorSkill(
        settings=settings,
        api_client=mock_api,
        formatter=OutputFormatter(settings),
        context_manager=ContextManager(settings),
        knowledge_base=KnowledgeBase(settings),
    )


class TestRfiValidation:
    """Test input validation for RFIs."""

    def test_validate_input_valid(self, seeded_db: Settings) -> None:
        skill = _make_skill(seeded_db)
        result = skill.validate_input(
            project="KRUPP-2026-TEST",
            issue="Structural drawings show W12x26 beams but specs call for W14x30.",
        )
        assert result["project_id"] == 1
        assert result["priority"] == "normal"
        assert result["to"] == "Architect"

    def test_validate_input_missing_project(self, seeded_db: Settings) -> None:
        skill = _make_skill(seeded_db)
        with pytest.raises(ValueError, match="Project code is required"):
            skill.validate_input(issue="Some issue description here.")

    def test_validate_input_missing_issue(self, seeded_db: Settings) -> None:
        skill = _make_skill(seeded_db)
        with pytest.raises(ValueError, match="Issue description is required"):
            skill.validate_input(project="KRUPP-2026-TEST")

    def test_validate_input_issue_too_short(self, seeded_db: Settings) -> None:
        skill = _make_skill(seeded_db)
        with pytest.raises(ValueError, match="at least 10 characters"):
            skill.validate_input(project="KRUPP-2026-TEST", issue="Short")

    def test_validate_input_invalid_priority(self, seeded_db: Settings) -> None:
        skill = _make_skill(seeded_db)
        with pytest.raises(ValueError, match="Priority must be"):
            skill.validate_input(
                project="KRUPP-2026-TEST",
                issue="Beam size discrepancy needs resolution.",
                priority="critical",
            )

    def test_validate_input_with_options(self, seeded_db: Settings) -> None:
        skill = _make_skill(seeded_db)
        result = skill.validate_input(
            project="KRUPP-2026-TEST",
            issue="Beam size discrepancy needs resolution.",
            to="Structural Engineer",
            priority="urgent",
            drawing_ref="S-301",
            spec_ref="05 12 00",
        )
        assert result["to"] == "Structural Engineer"
        assert result["priority"] == "urgent"
        assert result["drawing_ref"] == "S-301"
        assert result["spec_ref"] == "05 12 00"


class TestRfiPrompt:
    """Test prompt building for RFIs."""

    def test_build_prompt_structure(self, seeded_db: Settings) -> None:
        skill = _make_skill(seeded_db)
        context = skill.context_manager.load_by_code("KRUPP-2026-TEST")
        validated = {
            "project_id": 1,
            "project_code": "KRUPP-2026-TEST",
            "issue": "Beam size discrepancy.",
            "to": "Architect",
            "priority": "normal",
            "drawing_ref": None,
            "spec_ref": None,
        }
        messages = skill.build_prompt(validated, context)
        assert isinstance(messages, list)
        assert len(messages) >= 1
        assert messages[0]["role"] == "user"
        assert "ISSUE DESCRIPTION" in messages[0]["content"]

    def test_build_prompt_includes_context(self, seeded_db: Settings) -> None:
        skill = _make_skill(seeded_db)
        context = skill.context_manager.load_by_code("KRUPP-2026-TEST")
        validated = {
            "project_id": 1,
            "project_code": "KRUPP-2026-TEST",
            "issue": "Beam size discrepancy.",
            "to": "Architect",
            "priority": "high",
            "drawing_ref": "S-301",
            "spec_ref": None,
        }
        messages = skill.build_prompt(validated, context)
        content = messages[0]["content"]
        assert "Beam size discrepancy" in content
        assert "S-301" in content


class TestRfiOutput:
    """Test output formatting for RFIs."""

    def test_format_output_creates_file(self, seeded_db: Settings) -> None:
        skill = _make_skill(seeded_db)
        context = skill.context_manager.load_by_code("KRUPP-2026-TEST")
        validated = {
            "project_id": 1,
            "project_code": "KRUPP-2026-TEST",
            "issue": "Beam size discrepancy.",
            "to": "Architect",
            "priority": "normal",
        }
        output_path = skill.format_output(MOCK_RFI_JSON, context, validated)
        assert output_path.exists()
        assert output_path.suffix == ".docx"

    def test_format_output_content(self, seeded_db: Settings) -> None:
        skill = _make_skill(seeded_db)
        context = skill.context_manager.load_by_code("KRUPP-2026-TEST")
        validated = {
            "project_id": 1,
            "project_code": "KRUPP-2026-TEST",
            "issue": "Beam size discrepancy.",
            "to": "Architect",
            "priority": "normal",
        }
        output_path = skill.format_output(MOCK_RFI_JSON, context, validated)

        from docx import Document

        doc = Document(str(output_path))
        all_text = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    all_text.append(cell.text)
        full_text = "\n".join(all_text)
        assert "RFI" in full_text
        assert "KRUPP-2026-TEST" in full_text

    def test_format_output_persists_to_db(self, seeded_db: Settings) -> None:
        skill = _make_skill(seeded_db)
        context = skill.context_manager.load_by_code("KRUPP-2026-TEST")
        validated = {
            "project_id": 1,
            "project_code": "KRUPP-2026-TEST",
            "issue": "Beam size discrepancy.",
            "to": "Architect",
            "priority": "normal",
        }
        skill.format_output(MOCK_RFI_JSON, context, validated)

        with get_db(seeded_db) as conn:
            cursor = conn.execute("SELECT * FROM rfis WHERE project_id = 1")
            row = cursor.fetchone()
            assert row is not None
            assert row["rfi_number"] == 1
            assert row["status"] == "draft"
            assert "Structural Steel" in row["subject"]


class TestRfiEndToEnd:
    """Test full RFI lifecycle."""

    def test_execute_end_to_end(self, seeded_db: Settings) -> None:
        skill = _make_skill(seeded_db)
        result = skill.execute(
            project="KRUPP-2026-TEST",
            issue="Structural drawings show W12x26 beams but specs call for W14x30.",
        )
        assert result.output_path.exists()
        assert result.skill_name == "rfi_generator"
        assert result.project_code == "KRUPP-2026-TEST"

    def test_rfi_auto_numbering(self, seeded_db: Settings) -> None:
        """Two RFIs on same project get sequential numbers."""
        skill = _make_skill(seeded_db)
        skill.execute(
            project="KRUPP-2026-TEST",
            issue="First issue: beam size discrepancy at grid line C.",
        )
        skill.execute(
            project="KRUPP-2026-TEST",
            issue="Second issue: missing spec section for curtain wall system.",
        )

        with get_db(seeded_db) as conn:
            cursor = conn.execute(
                "SELECT rfi_number FROM rfis WHERE project_id = 1 "
                "ORDER BY rfi_number"
            )
            numbers = [row["rfi_number"] for row in cursor.fetchall()]
            assert numbers == [1, 2]
